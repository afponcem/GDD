#!/usr/bin/env python3
"""Parsea el Tablero de Indicadores GDD (xlsx) a un JSON consumido por el sitio.

Lee dos hojas con formato "jerárquico" idéntico:
  - "Resumen Semanal"   -> snapshot de la semana actual, por Territorio/Subgerencia
  - "Resumen Cump YTD"  -> snapshot acumulado del año, por Territorio/Subgerencia

Estructura de esas hojas (fija, ver .claude/skills/indicadores-sharepoint/SKILL.md
si esto cambia en el archivo fuente):
  fila 3: nombre de columna "sucio" (con sufijo de periodo, ej. "CAPTACIÓN TOTALS2")
  fila 4: nombre de columna canónico (ej. "CAPTACIÓN TOTAL")  <- se usa para la key
  fila 5: "Fecha Corte ->" + fecha por indicador
  fila 6: "Exigencia ->"   + meta por indicador
  fila 8: fila TOTAL RED ACHS (col A/B = código/nombre, resto = valores)
  marcador "TERRITORIO" en col B, luego filas de territorio (col A=código, col B=nombre)
  marcador "SUBGERENCIA" en col B, luego filas de subgerencia:
      col A = código, col B = nombre corto de territorio (solo en la 1a fila del grupo),
      col C = nombre de la subgerencia
  marcador "AGENCIAS" en col B, luego filas de agencia: col A/C = código/nombre,
      col B = nombre corto de territorio (solo en la 1a fila del grupo)
  marcador "JEFATURA" en col B, luego filas de persona. A diferencia de las
      demás secciones, acá el código/nombre de Subgerencia o Territorio que
      vuelve a aparecer (mismo código+nombre ya vistos) **cierra** el grupo
      de personas que viene ANTES, no abre el que sigue (ver el branch
      `elif section == "JEFATURA"` en parse_summary_sheet). El vínculo
      Jefatura -> Agencia no está en estas hojas en absoluto; se obtiene de
      la hoja cruda "YTD" (ver `build_jefatura_agencia_tags`).
"""
import json
import re
import sys
import unicodedata
from datetime import date, datetime, time, timezone
from pathlib import Path

import openpyxl

SHEET_CONFIGS = {
    "weekly": "Resumen Semanal",
    "ytd": "Resumen Cump YTD",
}
RAW_YTD_SHEET = "YTD"  # hoja cruda mensual, distinta de "Resumen Cump YTD"

# Col B en las filas de SUBGERENCIA/AGENCIAS solo trae el nombre corto del
# territorio (NORTE / METRO / SUR). Se mapea al código de territorio real.
TERRITORIO_SHORT_NAME_MAP = {
    "NORTE": "TERRITORIO NORTE",
    "METRO": "TERRITORIO METROPOLITANO",
    "SUR": "TERRITORIO SUR",
}

HEADER_ROW_DIRTY = 3
HEADER_ROW_CANON = 4
FECHA_CORTE_ROW = 5
EXIGENCIA_ROW = 6
DATA_START_ROW = 8
FIRST_INDICATOR_COL = 4  # columna D

# Columnas de la hoja cruda "YTD" usadas para el vínculo Jefatura -> Agencia.
RAW_YTD_CODE_COL = 4    # columna D: código de persona (o de agencia, en sus filas de rollup)
RAW_YTD_AGENCIA_COL = 6  # columna F: nombre de agencia (solo presente en alguna de las 3 filas REAL/META/CUMP de la persona)
RAW_YTD_LABEL_COL = 8    # columna H: "REAL" / "META" / "CUMP"


def slugify(text: str) -> str:
    text = unicodedata.normalize("NFKD", text).encode("ascii", "ignore").decode("ascii")
    text = text.strip().lower()
    text = re.sub(r"[^a-z0-9]+", "_", text)
    return text.strip("_")


def normalize_name(text) -> str:
    """Normaliza un nombre de agencia/entidad para poder compararlo entre
    hojas que lo escriben con o sin tilde (ej. 'PARQUE LAS AMÉRICAS' en la
    hoja cruda "YTD" vs 'PARQUE LAS AMERICAS' en Resumen Semanal)."""
    text = unicodedata.normalize("NFKD", str(text)).encode("ascii", "ignore").decode("ascii")
    return re.sub(r"\s+", " ", text.strip().upper())


def cell_to_jsonable(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, (date, time)):
        return value.isoformat()
    return value


def read_indicators(ws):
    """Devuelve {col_index: {"key":..., "label":..., "meta":..., "fecha_corte":...}}."""
    indicators = {}
    for col in range(FIRST_INDICATOR_COL, ws.max_column + 1):
        canon = ws.cell(row=HEADER_ROW_CANON, column=col).value
        if not canon:
            continue
        key = slugify(str(canon))
        if not key:
            continue
        indicators[col] = {
            "key": key,
            "label": str(canon).strip(),
            "meta": cell_to_jsonable(ws.cell(row=EXIGENCIA_ROW, column=col).value),
            "fecha_corte": cell_to_jsonable(ws.cell(row=FECHA_CORTE_ROW, column=col).value),
        }
    return indicators


def read_values(ws, row, indicators):
    values = {}
    for col, meta in indicators.items():
        v = ws.cell(row=row, column=col).value
        if v == "-" or v is None:
            continue
        values[meta["key"]] = cell_to_jsonable(v)
    return values


def build_jefatura_agencia_tags(wb):
    """Escanea la hoja cruda "YTD" (detalle mensual por persona, NO la hoja
    "Resumen Cump YTD") y devuelve {código_persona: nombre_agencia}.

    El nombre de agencia aparece en la columna F, pero no siempre en la
    misma de las 3 filas (REAL/META/CUMP) de esa persona — a veces en la
    primera, a veces en la tercera — así que se escanean las tres y se toma
    el primer valor no vacío encontrado.
    """
    if RAW_YTD_SHEET not in wb.sheetnames:
        return {}
    ws = wb[RAW_YTD_SHEET]
    tags = {}
    for r in range(1, ws.max_row + 1):
        label = ws.cell(row=r, column=RAW_YTD_LABEL_COL).value
        if label not in ("REAL", "META", "CUMP"):
            continue
        code = ws.cell(row=r, column=RAW_YTD_CODE_COL).value
        tag = ws.cell(row=r, column=RAW_YTD_AGENCIA_COL).value
        if not code or not tag:
            continue
        tags.setdefault(str(code).strip(), str(tag).strip())
    return tags


def parse_summary_sheet(ws):
    """Parsea una hoja "Resumen ..." a (tree, indicator_catalog, agencias_by_name).

    `tree` queda anidado Total -> Territorio -> Subgerencia -> Jefatura
    (Agencia se inserta después, en `nest_agencias`, porque requiere cruzar
    con la hoja cruda "YTD"). `agencias_by_name` son los nodos de la sección
    AGENCIAS, sin colgar todavía de ningún padre, indexados por
    `normalize_name(nombre)` -> (nodo, territorio_padre).
    """
    indicators = read_indicators(ws)

    total_code = str(ws.cell(row=DATA_START_ROW, column=1).value or "TOTAL_RED_ACHS").strip()
    total_name = str(ws.cell(row=DATA_START_ROW, column=2).value or "TOTAL RED ACHS").strip()

    tree = {
        "code": total_code,
        "name": total_name,
        "level": "total",
        "values": read_values(ws, DATA_START_ROW, indicators),
        "children": [],
    }
    territorios_by_short_name = {}
    territorios_by_code = {}
    subgerencias_by_code = {}
    agencias_by_name = {}

    def resolve_territorio(short_name):
        full_name = TERRITORIO_SHORT_NAME_MAP.get(short_name, short_name)
        return territorios_by_short_name.get(full_name)

    section = None
    current_territorio_short = None
    jefatura_buffer = []  # personas acumuladas hasta el próximo marcador de cierre
    row = DATA_START_ROW + 1
    while row <= ws.max_row:
        col_a = ws.cell(row=row, column=1).value
        col_b = ws.cell(row=row, column=2).value
        col_c = ws.cell(row=row, column=3).value

        if col_a is None and col_b in ("TERRITORIO", "SUBGERENCIA", "AGENCIAS", "JEFATURA"):
            section = col_b
            current_territorio_short = None
            row += 1
            continue

        if section == "TERRITORIO" and col_a:
            node = {
                "code": str(col_a).strip(),
                "name": str(col_b).strip() if col_b else str(col_a).strip(),
                "level": "territorio",
                "values": read_values(ws, row, indicators),
                "children": [],
            }
            tree["children"].append(node)
            territorios_by_short_name[node["name"]] = node
            territorios_by_code[node["code"]] = node

        elif section == "SUBGERENCIA" and col_a:
            if col_b:
                current_territorio_short = str(col_b).strip()
            parent = resolve_territorio(current_territorio_short)
            node = {
                "code": str(col_a).strip(),
                "name": str(col_c).strip() if col_c else str(col_a).strip(),
                "level": "subgerencia",
                "values": read_values(ws, row, indicators),
                "children": [],
            }
            subgerencias_by_code[node["code"]] = node
            if parent is not None:
                parent["children"].append(node)
            else:
                tree["children"].append(node)  # fallback si no calzó el territorio

        elif section == "AGENCIAS" and col_a:
            if col_b:
                current_territorio_short = str(col_b).strip()
            parent = resolve_territorio(current_territorio_short)
            name = str(col_c).strip() if col_c else str(col_a).strip()
            node = {
                "code": str(col_a).strip(),
                "name": name,
                "level": "agencia",
                "values": read_values(ws, row, indicators),
                "children": [],
            }
            # No se cuelga de inmediato: `nest_agencias` decide si va bajo la
            # Subgerencia (cruzando con la hoja cruda "YTD") o, si no matchea
            # con ninguna jefatura, cae aquí como respaldo bajo el territorio.
            agencias_by_name[normalize_name(name)] = (node, parent or tree)

        elif section == "JEFATURA" and col_a:
            code = str(col_a).strip()
            name = str(col_c).strip() if col_c else code
            candidate_subgerencia = subgerencias_by_code.get(code)
            candidate_territorio = territorios_by_code.get(code)
            if candidate_subgerencia is not None and candidate_subgerencia["name"] == name:
                # Esta fila repite el código/nombre de una Subgerencia: cierra
                # (no abre) el grupo de personas acumulado desde el último
                # marcador — son las jefaturas de esa subgerencia.
                candidate_subgerencia["children"].extend(jefatura_buffer)
                jefatura_buffer = []
            elif candidate_territorio is not None and candidate_territorio["name"] == name:
                # Mismo patrón pero a nivel Territorio: cargos que no caen
                # bajo ninguna Subgerencia específica (liderazgo territorial).
                candidate_territorio["children"].extend(jefatura_buffer)
                jefatura_buffer = []
            else:
                jefatura_buffer.append({
                    "code": code,
                    "name": name,
                    "level": "jefatura",
                    "values": read_values(ws, row, indicators),
                    "children": [],
                })

        row += 1

    if jefatura_buffer:
        # No debería quedar nada sin cerrar (la hoja siempre termina con un
        # marcador de Territorio) — si pasa, mejor no perder los datos.
        tree["children"].extend(jefatura_buffer)

    indicator_catalog = {
        meta["key"]: {"label": meta["label"], "meta": meta["meta"], "fecha_corte": meta["fecha_corte"]}
        for meta in indicators.values()
    }
    return tree, indicator_catalog, agencias_by_name


def nest_agencias(tree, jefatura_tags, agencias_by_name):
    """Inserta el nivel Agencia entre Subgerencia y Jefatura, agrupando las
    jefaturas de cada subgerencia según `jefatura_tags` (código -> nombre de
    agencia, desde la hoja cruda "YTD"). Cuando el nombre de agencia calza
    con uno ya visto en la sección AGENCIAS de esta misma hoja, reusa ese
    nodo (con sus valores reales); si no calza (ej. una jefatura a cargo de
    una zona combinada tipo "ARICA-IQUIQUE" que no es una agencia individual
    de la sección AGENCIAS), crea un nodo liviano sin valores propios — la
    fila queda igual disponible para filtrar/drill-down, solo sin % de
    cumplimiento agregado a ese nivel.

    Las agencias que no calzan con ninguna jefatura (o cuyo territorio no
    tiene jefaturas parseadas) quedan como respaldo colgando directo del
    territorio, igual que antes de tener este cruce.
    """
    consumed = set()

    def group_jefaturas_by_agencia(container_node):
        """Agrupa los hijos nivel 'jefatura' de `container_node` por su tag
        de agencia; deja cualquier otro hijo (subgerencias, jefaturas sin
        tag) tal cual. Se usa tanto para los hijos de una Subgerencia como
        para los hijos directos de un Territorio (cargos que no cuelgan de
        ninguna subgerencia pero sí tienen agencia asignada en la hoja
        cruda "YTD")."""
        groups = {}
        passthrough = []
        for child in container_node["children"]:
            tag = jefatura_tags.get(child["code"]) if child["level"] == "jefatura" else None
            if not tag:
                passthrough.append(child)
                continue
            groups.setdefault(tag, []).append(child)

        new_children = passthrough
        for tag, members in groups.items():
            norm = normalize_name(tag)
            existing = agencias_by_name.get(norm)
            if existing:
                agencia_node, _parent = existing
                agencia_node["children"] = members
                new_children.append(agencia_node)
                consumed.add(norm)
            else:
                new_children.append({
                    "code": tag,
                    "name": tag,
                    "level": "agencia",
                    "values": {},
                    "children": members,
                })
        container_node["children"] = new_children

    for territorio_node in tree["children"]:
        if territorio_node["level"] != "territorio":
            continue
        for child in territorio_node["children"]:
            if child["level"] == "subgerencia":
                group_jefaturas_by_agencia(child)
        # Cargos que cuelgan directo del territorio (sin subgerencia) también
        # pueden tener agencia asignada en la hoja cruda "YTD".
        group_jefaturas_by_agencia(territorio_node)

    for norm, (agencia_node, parent_node) in agencias_by_name.items():
        if norm not in consumed:
            parent_node["children"].append(agencia_node)


def parse_workbook(xlsx_path: Path) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    jefatura_tags = build_jefatura_agencia_tags(wb)

    result = {
        "generated_at": datetime.now(timezone.utc).isoformat(),
        "source_file": xlsx_path.name,
    }
    for out_key, sheet_name in SHEET_CONFIGS.items():
        if sheet_name not in wb.sheetnames:
            raise KeyError(f"No se encontró la hoja '{sheet_name}' en {xlsx_path}")
        tree, indicators, agencias_by_name = parse_summary_sheet(wb[sheet_name])
        nest_agencias(tree, jefatura_tags, agencias_by_name)
        result[out_key] = {"hierarchy": tree, "indicators": indicators}

    return result


def main():
    if len(sys.argv) != 3:
        print(f"Uso: {sys.argv[0]} <archivo.xlsx> <salida.json>", file=sys.stderr)
        sys.exit(1)

    xlsx_path = Path(sys.argv[1])
    out_path = Path(sys.argv[2])

    data = parse_workbook(xlsx_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    out_path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"OK: {out_path} ({out_path.stat().st_size} bytes)")


if __name__ == "__main__":
    main()
