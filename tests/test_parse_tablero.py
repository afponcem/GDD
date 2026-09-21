"""Test del parser contra un workbook sintético que reproduce el layout real
de las hojas "Resumen Semanal" / "Resumen Cump YTD" (ver
.claude/skills/indicadores-sharepoint/SKILL.md para el detalle del formato).

No depende del archivo real de ACHS ni de credenciales de SharePoint —
corre en CI en cada push.
"""
import sys
from pathlib import Path

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
from parse_tablero import parse_summary_sheet, slugify  # noqa: E402


def build_fixture_sheet(wb, name):
    ws = wb.create_sheet(name)

    # fila 3/4: nombre sucio / nombre canónico
    ws.cell(row=3, column=4, value="CAPTACIÓN TOTALS2")
    ws.cell(row=4, column=4, value="CAPTACIÓN TOTAL")
    ws.cell(row=3, column=5, value="Accidentes (vista de gestión)S2")
    ws.cell(row=4, column=5, value="Accidentes (vista de gestión)")

    # fila 5/6: fecha de corte / exigencia
    ws.cell(row=5, column=4, value="2026-01-05")
    ws.cell(row=6, column=4, value=1)
    ws.cell(row=5, column=5, value="2026-01-05")
    ws.cell(row=6, column=5, value="(Valor Real)")

    # fila 8: TOTAL
    ws.cell(row=8, column=1, value="TOTAL RED ACHS")
    ws.cell(row=8, column=2, value="TOTAL RED ACHS")
    ws.cell(row=8, column=4, value=0.95)
    ws.cell(row=8, column=5, value=10)

    # marcador TERRITORIO + 2 territorios
    ws.cell(row=10, column=2, value="TERRITORIO")
    ws.cell(row=11, column=1, value="GRTR1020")
    ws.cell(row=11, column=2, value="TERRITORIO NORTE")
    ws.cell(row=11, column=4, value=1.1)
    ws.cell(row=11, column=5, value=4)

    ws.cell(row=12, column=1, value="GRTR1010")
    ws.cell(row=12, column=2, value="TERRITORIO METROPOLITANO")
    ws.cell(row=12, column=4, value=1.2)
    ws.cell(row=12, column=5, value=5)

    # marcador SUBGERENCIA + 2 subgerencias bajo NORTE, 1 bajo METRO
    ws.cell(row=14, column=2, value="SUBGERENCIA")
    ws.cell(row=15, column=1, value="SGRG0003")
    ws.cell(row=15, column=2, value="NORTE")
    ws.cell(row=15, column=3, value="ARICA Y TARAPACA")
    ws.cell(row=15, column=4, value=0.8)
    ws.cell(row=15, column=5, value=2)

    ws.cell(row=16, column=1, value="SGRG0001")
    # col B vacía: forward-fill del territorio corto "NORTE"
    ws.cell(row=16, column=3, value="ANTOFAGASTA")
    ws.cell(row=16, column=4, value="-")  # sin dato -> se omite
    ws.cell(row=16, column=5, value=0)

    ws.cell(row=17, column=1, value="SGRG0011")
    ws.cell(row=17, column=2, value="METRO")
    ws.cell(row=17, column=3, value="METRO NORTE")
    ws.cell(row=17, column=4, value=0.9)
    ws.cell(row=17, column=5, value=1)

    # marcador AGENCIAS: cuelga directo del territorio (no de la subgerencia,
    # la planilla no da ese vínculo)
    ws.cell(row=19, column=2, value="AGENCIAS")
    ws.cell(row=20, column=1, value="IQUIQUE")
    ws.cell(row=20, column=2, value="NORTE")
    ws.cell(row=20, column=3, value="IQUIQUE")
    ws.cell(row=20, column=4, value=0.99)

    # marcador JEFATURA: 1 cargo a nivel territorio NORTE (antes de cualquier
    # marcador de subgerencia) + 1 fila que repite la subgerencia SGRG0003
    # como encabezado + 1 persona bajo esa subgerencia + 1 cargo a nivel
    # territorio METRO (regresión: no debe heredar la subgerencia de NORTE)
    ws.cell(row=22, column=2, value="JEFATURA")
    ws.cell(row=23, column=1, value="JECPX001")
    ws.cell(row=23, column=2, value="NORTE")
    ws.cell(row=23, column=3, value="JEFE TERRITORIAL EJEMPLO")
    ws.cell(row=23, column=4, value=0.5)

    ws.cell(row=24, column=1, value="SGRG0003")
    ws.cell(row=24, column=3, value="ARICA Y TARAPACA")  # debe calzar con sub1["name"]
    ws.cell(row=24, column=4, value=0.8)  # repite el valor de la subgerencia

    ws.cell(row=25, column=1, value="JECPX002")
    ws.cell(row=25, column=3, value="PERSONA EJEMPLO")
    ws.cell(row=25, column=4, value=0.7)

    ws.cell(row=26, column=1, value="JECPX003")
    ws.cell(row=26, column=2, value="METRO")
    ws.cell(row=26, column=3, value="JEFE METRO EJEMPLO")
    ws.cell(row=26, column=4, value=0.6)

    return ws


def test_parses_hierarchy_and_values():
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ws = build_fixture_sheet(wb, "Resumen Semanal")

    tree, indicators = parse_summary_sheet(ws)

    assert indicators["captacion_total"]["label"] == "CAPTACIÓN TOTAL"
    assert indicators["captacion_total"]["meta"] == 1
    assert indicators["accidentes_vista_de_gestion"]["meta"] == "(Valor Real)"

    assert tree["code"] == "TOTAL RED ACHS"
    assert tree["values"]["captacion_total"] == 0.95

    assert len(tree["children"]) == 2
    territorio, territorio_metro = tree["children"]
    assert territorio["code"] == "GRTR1020"
    assert territorio["level"] == "territorio"
    assert territorio_metro["code"] == "GRTR1010"

    subgerencias = [c for c in territorio["children"] if c["level"] == "subgerencia"]
    assert len(subgerencias) == 2
    sub1, sub2 = subgerencias
    assert sub1["code"] == "SGRG0003"
    assert sub1["name"] == "ARICA Y TARAPACA"
    assert sub1["values"]["captacion_total"] == 0.8

    assert sub2["code"] == "SGRG0001"
    assert sub2["name"] == "ANTOFAGASTA"
    # "-" se omite del dict de valores (no se guarda como 0 ni como string)
    assert "captacion_total" not in sub2["values"]
    assert sub2["values"]["accidentes_vista_de_gestion"] == 0

    # Agencia cuelga directo del territorio (no de la subgerencia)
    agencias = [c for c in territorio["children"] if c["level"] == "agencia"]
    assert len(agencias) == 1
    assert agencias[0]["code"] == "IQUIQUE"
    assert agencias[0]["values"]["captacion_total"] == 0.99

    # Jefatura territorial (antes del primer marcador de subgerencia) cuelga
    # directo del territorio
    jefaturas_territorio = [c for c in territorio["children"] if c["level"] == "jefatura"]
    assert len(jefaturas_territorio) == 1
    assert jefaturas_territorio[0]["code"] == "JECPX001"

    # La fila que repite código+nombre de SGRG0003 es un encabezado de grupo,
    # no una persona: no debe aparecer como nodo "jefatura", y la persona
    # que sigue debe colgar de esa subgerencia, no del territorio.
    all_jefatura_codes = [
        c["code"]
        for t in tree["children"]
        for c in t["children"]
        if c["level"] == "jefatura"
    ] + [
        p["code"]
        for t in tree["children"]
        for s in t["children"]
        if s["level"] == "subgerencia"
        for p in s["children"]
    ]
    assert "SGRG0003" not in all_jefatura_codes
    assert sub1["children"][0]["code"] == "JECPX002"
    assert sub1["children"][0]["level"] == "jefatura"
    assert sub1["children"][0]["values"]["captacion_total"] == 0.7

    # Regresión: un cargo territorial de METRO (sin marcador de subgerencia
    # propio) no debe heredar la subgerencia SGRG0003 dejada por NORTE.
    jefaturas_metro = [c for c in territorio_metro["children"] if c["level"] == "jefatura"]
    assert len(jefaturas_metro) == 1
    assert jefaturas_metro[0]["code"] == "JECPX003"
    subgerencia_metro = [c for c in territorio_metro["children"] if c["level"] == "subgerencia"][0]
    assert subgerencia_metro["children"] == []


def test_slugify_handles_accents_and_punctuation():
    assert slugify("CAPTACIÓN TOTAL") == "captacion_total"
    assert slugify("Accidentes (vista de gestión)") == "accidentes_vista_de_gestion"
    assert slugify("  Cobertura RC (M+)  ") == "cobertura_rc_m"


if __name__ == "__main__":
    test_parses_hierarchy_and_values()
    test_slugify_handles_accents_and_punctuation()
    print("OK: todos los tests pasaron")
